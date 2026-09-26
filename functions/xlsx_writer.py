"""Write analysis summary tables to xlsx files."""

from pathlib import Path

import polars as pl
from openpyxl import Workbook, load_workbook


def write_cell_summary_xlsx(cell_df: pl.DataFrame, output_path: Path) -> None:
    """Write a count_unique_cells() result to output_path as a single-sheet xlsx.

    Filenames lists are joined into comma-separated strings, since xlsx cells
    can't hold a real list.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Cells"
    ws.append(["ANIMAL_ID", "SLICE", "AT", "Filenames", "n_images"])
    for row in cell_df.iter_rows(named=True):
        ws.append([row["ANIMAL_ID"], row["SLICE"], row["AT"], ", ".join(row["Filenames"]), row["n_images"]])
    wb.save(output_path)


def write_stats_xlsx(tables: dict[str, list[dict]], output_path: Path) -> None:
    """Add one sheet per table ({sheet name: rows of {column: value}}) to output_path, replacing same-name sheets.

    Other sheets (e.g. "Cells") are kept; a missing file is created. An empty table gets a "(none)" row.
    """
    exists = output_path.exists()
    wb = load_workbook(output_path) if exists else Workbook()
    if not exists:
        wb.remove(wb.active)  # drop the default empty sheet
    for name, rows in tables.items():
        if name in wb.sheetnames:
            wb.remove(wb[name])
        ws = wb.create_sheet(name)
        if not rows:
            ws.append(["(none)"])
            continue
        ws.append(list(rows[0].keys()))
        for row in rows:
            ws.append(list(row.values()))
    wb.save(output_path)
